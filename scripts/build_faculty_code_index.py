from pathlib import Path
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'Team_17_Faculty_Code_Index.docx'
GROUPS=[('Module 1 Preprocessing and validation',[
 ('CSV reading and validation','hotel_compiler/csv_input.py','Headers, rows, empty reviews, duplicate IDs and limits.'),
 ('Compiler front end','hotel_compiler/front_end.py','Validation followed by lexical, syntax and semantic processing.')]),
 ('Module 2 Lexical and grammatical analysis',[
 ('SentiWordNet lexical analysis','hotel_compiler/lexer.py','Tokenization, longest phrase matching, contractions and guarded morphology.'),
 ('Controlled syntax grammar','hotel_compiler/parser.py','Predictive parsing and supported clause structures.'),
 ('Additional grammar and recovery','hotel_compiler/expanded_parser.py','Bounded recovery, coordination and limited reference handling.'),
 ('Hybrid grammatical analysis','hotel_compiler/hybrid_nlp.py','analyze_language: Pattern POS/chunks, source spans and opinion attachment.')]),
 ('Module 3 Semantic attachment and lexical evidence',[
 ('Categories, issues and compatibility','hotel_compiler/semantics.py','analyze_semantics: bind opinions to targets and check semantic guards.'),
 ('SentiWordNet vocabulary','hotel_compiler/resources/domain_expanded.json','Configured aspects, opinions, aliases, senses and modifiers.'),
 ('Native SentiWordNet values','hotel_compiler/resources/sentiwordnet_expanded.tsv','Actual packaged synset rows with positive and negative source values.'),
 ('Hybrid tools and fusion','hotel_compiler/hybrid.py','HybridEngine._score_native and score_text: VADER, TextBlob, evidence flags and fusion.'),
 ('Guarded vocabulary adaptation','hotel_compiler/hybrid_vocabulary.py','normalize_opinions: scoring-copy substitutions and retained evidence.'),
 ('Hybrid lexical catalog','hotel_compiler/resources/hybrid_vocabulary.json','Reviewed lexical forms and category guards.')]),
 ('Module 4 Adjustments and final records',[
 ('SentiWordNet final scoring','hotel_compiler/sentiwordnet.py','evaluate: native P-N, domain adjustments, degree, negation and clipping.'),
 ('Hybrid context adjustments','hotel_compiler/hybrid_nlp.py','domain_adjustment: match hotel rules and handle opposing rule signs.'),
 ('Hybrid aspect and rule configuration','hotel_compiler/resources/hybrid_domain.json','Aspect dictionaries, context patterns and assigned rule scores.'),
 ('Five labels and Mixed sentiment','hotel_compiler/sentiment_policy.py','SentimentPolicy and overall_polarity: thresholds and coexistence of polarities.'),
 ('SentiWordNet orchestration and status','hotel_compiler/pipeline.py','analyze_review: connect stages and set Complete, Partial or Unsupported.'),
 ('Review score and intermediate records','hotel_compiler/models.py','AnalysisResult.review_score, source-linked records and evidence types.'),
 ('Hybrid synthesis and status','hotel_compiler/hybrid.py','analyze_hybrid_review and HybridAnalysis: aspect mean, Mixed and diagnostics.')]),
 ('Evaluation datasets and demonstration',[
 ('SemEval human annotation reader','hotel_compiler/research_data.py','load_semeval: read original sentences and human opinions from XML.'),
 ('Tuple generation and matching','hotel_compiler/research_runner.py','evaluate_semeval: mapped gold/predicted sets, intersections and tuple scores.'),
 ('Classification metrics and mapping','hotel_compiler/research_metrics.py','classification_metrics, category crosswalk and cluster bootstrap intervals.'),
 ('Real dataset cleaning','scripts/prepare_clean_booking_demo.py','Marker/whitespace cleanup, deduplication, sampling and provenance.'),
 ('Website backend','hotel_compiler/web.py','HTTP routes, review requests and CSV batch handling.'),
 ('Dashboard and upload interface','hotel_compiler/web_assets/dashboard.js','Automatic headers, CSV upload, charts and individual-review display.')])]

def hyperlink(p,path):
    target=ROOT/path;assert target.is_file(),target
    rel=p.part.relate_to(target.as_uri(),RT.HYPERLINK,is_external=True)
    link=OxmlElement('w:hyperlink');link.set(qn('r:id'),rel)
    run=OxmlElement('w:r');pr=OxmlElement('w:rPr')
    color=OxmlElement('w:color');color.set(qn('w:val'),'1155CC');pr.append(color)
    u=OxmlElement('w:u');u.set(qn('w:val'),'single');pr.append(u)
    run.append(pr);t=OxmlElement('w:t');t.text=Path(path).name;run.append(t);link.append(run);p._p.append(link)

doc=Document();s=doc.sections[0]
s.page_width=Inches(8.5);s.page_height=Inches(11)
s.top_margin=s.bottom_margin=Inches(.65);s.left_margin=s.right_margin=Inches(.65)
for name in ['Normal','Title','Heading 1']:
    st=doc.styles[name];st.font.name='Calibri';st.font.color.rgb=RGBColor(0,0,0)
doc.styles['Normal'].font.size=Pt(10)
doc.styles['Normal'].paragraph_format.space_after=Pt(3)
doc.styles['Title'].font.size=Pt(22)
doc.styles['Heading 1'].font.size=Pt(12)
doc.styles['Heading 1'].paragraph_format.space_before=Pt(9)
doc.add_paragraph('Faculty Viva Code Reference','Title')
doc.add_paragraph('Team 17  |  Tanay Mohta and Aditya Rawat')
doc.add_paragraph('Use this index to open the implementation relevant to a faculty question. In desktop Word, hold Ctrl and click a blue filename. The links target verified files in this project on this computer.')
doc.add_paragraph('Keep these five files ready first: pipeline.py, semantics.py, sentiwordnet.py, hybrid.py and research_runner.py. Python files apply logic; JSON/TSV files contain the configured rules and lexical data.')
for title,rows in GROUPS:
    doc.add_paragraph(title,'Heading 1')
    table=doc.add_table(rows=1,cols=3);table.autofit=False
    widths=[2.0,1.8,3.4]
    for c,w in zip(table.columns,widths):c.width=Inches(w)
    for c,t in zip(table.rows[0].cells,['Faculty question or purpose','Open file','What to show']):c.text=t
    for purpose,path,detail in rows:
        cells=table.add_row().cells;cells[0].text=purpose;hyperlink(cells[1].paragraphs[0],path);cells[2].text=detail
    for ri,row in enumerate(table.rows):
        row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
        if ri==0:row._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
        for cell,w in zip(row.cells,widths):
            cell.width=Inches(w)
            if ri==0:
                shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'EEEEEE');cell._tc.get_or_add_tcPr().append(shade)
            for p in cell.paragraphs:
                p.paragraph_format.space_after=Pt(4);p.paragraph_format.space_before=Pt(3)
                for r in p.runs:r.bold=ri==0
    borders=OxmlElement('w:tblBorders')
    for side in ['top','left','bottom','right','insideH','insideV']:
        e=OxmlElement('w:'+side);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
    table._tbl.tblPr.append(borders)
doc.add_paragraph('The current paper uses two systems only: SentiWordNet and Hybrid. The filename domain_expanded.json identifies the current SentiWordNet configuration; it is not a third system. The hybrid independently uses Pattern grammar and its own attachment rules, rather than the SentiWordNet AST.')
for e in doc.styles.element.xpath('.//w:pBdr')+doc.element.xpath('.//w:pBdr'):e.getparent().remove(e)
doc.save(OUT)
links=[r for r in Document(OUT).part.rels.values() if r.reltype==RT.HYPERLINK]
assert len(links)>=20
print(f'{OUT}\nVerified {len(links)} hyperlink relationships and all target files.')
