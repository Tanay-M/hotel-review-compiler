"""Create the short post-14-October implementation report."""
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'output/reports/Team_17_Implementation_Update_After_14_October.docx'
doc=Document()
sec=doc.sections[0]
sec.page_width=Inches(8.5);sec.page_height=Inches(11)
sec.top_margin=sec.bottom_margin=Inches(.7)
sec.left_margin=sec.right_margin=Inches(.8)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','List Bullet']:
    st=doc.styles[name];st.font.name='Calibri';st.font.color.rgb=RGBColor(0,0,0)
    st.paragraph_format.space_after=Pt(5)
doc.styles['Normal'].font.size=Pt(11)
doc.styles['Normal'].paragraph_format.line_spacing=1.04
doc.styles['Title'].font.size=Pt(20)
doc.styles['Title'].paragraph_format.space_after=Pt(5)
doc.styles['Heading 1'].font.size=Pt(12)
doc.styles['Heading 1'].font.bold=True
doc.styles['Heading 1'].paragraph_format.space_before=Pt(10)
doc.styles['Heading 1'].paragraph_format.space_after=Pt(4)
doc.styles['Heading 1'].paragraph_format.keep_with_next=True

def p(text):doc.add_paragraph(text)
def h(text):doc.add_paragraph(text,'Heading 1')
def bullets(items):
    for text in items:doc.add_paragraph(text,'List Bullet')

doc.add_paragraph('Hotel Review Analysis Implementation Update','Title')
p('Team 17  |  Tanay Mohta and Aditya Rawat')
p('Reporting period: 15 to 21 October 2026')
p('This report summarizes the work implemented after the 14 October project stage. We added an independent VADER–TextBlob hybrid, broadened the SentiWordNet approach, and extended real-data evaluation and demonstration support. The current comparison contains two approaches: SentiWordNet and Hybrid.')

h('1 Independent hybrid architecture')
p('We implemented a second analysis pipeline that receives the original review independently. It does not depend on acceptance by the SentiWordNet lexer, parser or opinion vocabulary. Both approaches retain the six categories Room, Staff, Service, Food, Cleanliness and Facilities.')
bullets([
    'TextBlob PatternParser supplies part-of-speech tags and shallow phrase chunks. Our rules identify aspect mentions, divide relevant clauses and attach opinions to their targets.',
    'Attachment supports adjective–noun expressions, subject–predicate opinions, selected opinion verbs, coordinated targets and bounded references to an unambiguous previous target.',
    'Each extracted context is scored by VADER and TextBlob PatternAnalyzer. Whole-review lexical sentiment is kept separately from aspect sentiment.'
])

h('2 Hybrid scoring and traceability')
bullets([
    'VADER contributes compound sentiment and positive, neutral and negative proportions. TextBlob contributes polarity, subjectivity and assessed expressions; subjectivity is not confidence.',
    'When both tools provide lexical evidence, their scores use equal weights. When only one provides evidence, its score is retained. Missing lexical evidence is distinguished from supported Neutral sentiment; component failures remain errors.',
    'Explicit hotel-context rules handle conditions such as quiet rooms and equipment failure. The system retains component scores, fusion, matched rule and final score. A rule may supply sentiment even when lexical fusion is unavailable.',
    'Five descriptive bands were introduced: Strong Negative, Weak Negative, Neutral, Weak Positive and Strong Positive. Mixed remains available when positive and negative opinions coexist.'
])

h('3 Vocabulary and grammatical coverage')
p('The hybrid vocabulary was broadened using inspected Booking Italy reviews. It now contains 671 aspect phrases and 497 reviewed lexical forms alongside the native tool dictionaries. Eligible unknown forms can receive category-approved substitutes in a scoring copy; original review wording, source spans, substitutions and original/effective scores remain visible.')
p('We also broadened SentiWordNet for a more substantial comparison. Its current resources contain 772 aspect phrases, 21,181 opinion entries, 752 aliases and 14,311 genuine source synsets. Native positive and negative source values remain unchanged. Configured hotel senses take precedence; additional native entries use a first lemma-ranked sense policy.')

doc.add_page_break()
h('4 SentiWordNet robustness improvements')
p('The deterministic compiler gained bounded recovery for opinions inside otherwise unsupported clauses, selected verb–object opinions, guarded comparative/superlative handling, coordinated predicates and targets, contraction handling and limited explicit-reference inheritance. Explicit absence such as “no hot water” is represented with a traceable availability interpretation. Guards prevent false stemming and distinguish “not only” from ordinary negation. Unknown or ambiguous language retains diagnostics instead of being silently accepted.')

h('5 Real datasets and reproducible evaluation')
p('We prepared a cleaned derivative of the real Crawlfeeds Booking Italy export. From 4,367 source rows, documented cleaning removes 1,266 empty bodies and 12 exact cleaned duplicates, retaining 3,089 reviews. Cleaning removes missing markers and nonprinting controls and normalizes whitespace; it does not paraphrase reviews, correct sentiment wording or select successful predictions.')
p('A uniform random 500-review demo subset was generated with a fixed seed, without analyzer-based selection. Input CSVs contain only a serial number and review. Source maps, cleaning records and hashes retain the derivation. The full corpus is supplied in four parts to respect the upload limit. The 500 reviews are a subset of the 3,089, not additional independent data.')
p('The offline runner compares both approaches on identical text, preserves individual outcomes and exports review tables, opinion records, diagnostics and evidence. It reports annotated SemEval classification and category–polarity metrics separately from coverage on unannotated Booking reviews. Missing and Mixed predictions remain errors in the single-polarity benchmark task.')

h('6 Demonstration interface and reporting')
p('The existing WebUI was extended to expose Hybrid analysis, same-review comparison, hybrid CSV processing and individual result inspection. Trace views show grammatical information, attached aspects, component scores, normalization, domain rules and final records. Dashboards and exports identify the actual analysis method. Start_Demo.bat and Stop_Demo.bat support local demonstration.')
p('The IEEE paper was revised throughout, including its abstract, four-module methodology, pseudocode, architecture diagrams, tables, graphs and conclusion. It now presents only SentiWordNet versus Hybrid using the latest results. Reproducible aggregate figures and tables, an implementation reference and a detailed viva guide were also prepared. The current regression suite passes 210 authored software tests; these verify program contracts rather than real-world sentiment accuracy.')

h('7 Current project status')
p('On 148 mapped SemEval development cases, SentiWordNet achieves 48.65% accuracy and Hybrid 57.43%. Hybrid has higher tuple-F1 (0.6450 versus 0.5912), while SentiWordNet has slightly higher macro-F1 (0.4037 versus 0.3999). On the 3,089 cleaned Booking reviews, at least one aspect is scored for 81.42% and 87.41%, respectively; strict completeness remains 1.65% and 7.96%. Booking coverage is not accuracy, and the inspected benchmark is not an untouched test set.')
p('Both approaches run locally without an LLM call or database. Complex language and out-of-scope topics remain limitations. Fusion/override ablations, fresh holdout annotation and independently validated five-class strength accuracy remain future work.')

footer=sec.footer.paragraphs[0]
footer.alignment=2
r=footer.add_run('Team 17  |  Implementation update  |  ');r.font.size=Pt(9)
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
for element in doc.styles.element.xpath('.//w:pBdr') + doc.element.xpath('.//w:pBdr'):
    element.getparent().remove(element)
OUT.parent.mkdir(parents=True,exist_ok=True)
doc.save(OUT)
print(OUT)
print('Words:',sum(len(p.text.split()) for p in doc.paragraphs))
