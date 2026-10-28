"""Structural and measured-value audit, supplementary to visual page review."""
from pathlib import Path
from zipfile import ZipFile,ZIP_DEFLATED
from hashlib import sha256
import csv
import json
import re
from docx import Document
from docx.oxml.ns import qn
import pdfplumber

ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'output/paper'
RUN=ROOT/'output/experiments/2026-10-21/expanded-sentiwordnet-release'
FIGURES=['architecture_approved','performance_comparison','dl_joint_pair_f1']
DL=ROOT/'output/neural/2026-10-23-dl-pair-evaluation'
DL_RUN=ROOT/'output/neural/2026-10-23-dl-pair-human-v1'


def load(path):
    with path.open(encoding='utf-8-sig',newline='') as h:return list(csv.DictReader(h))


def main():
    path=OUT/'Team_17_IEEE_Conference_Paper.docx';doc=Document(path)
    paragraphs=[p.text for p in doc.paragraphs];text='\n'.join(paragraphs)
    reference_lines=[p for p in paragraphs if re.match(r'^\[\d+\]  ',p)]
    assert len(reference_lines)==35
    assert [int(re.match(r'\[(\d+)\]',p).group(1)) for p in reference_lines]==list(range(1,36))
    body=text.split('REFERENCES')[0]
    first=[]
    for value in re.findall(r'\[(\d+)\]',body):
        if int(value) not in first:first.append(int(value))
    assert first==list(range(1,36)),first
    assert len(doc.inline_shapes)==3
    assert len(doc.tables)==5
    assert sum(p.startswith('Algorithm ') for p in paragraphs)==5
    assert sum(p.startswith('Input:') for p in paragraphs)==5
    assert sum(p.startswith('Output:') for p in paragraphs)==5
    from build_conference_paper import ALGORITHMS
    for number, (title, lines) in ALGORITHMS.items():
        start=paragraphs.index(f'Algorithm {number}  {title}')
        actual=paragraphs[start+1:start+1+len(lines)]
        assert len(actual)==len(lines)
        steps=[p for p in actual if re.match(r'^\d+:\t',p)]
        assert len(steps)==len(lines)-2
        assert [int(p.split(':',1)[0]) for p in steps]==list(range(1,len(steps)+1))
    assert [int(re.match(r'Fig\. (\d+)\.',p)[1]) for p in paragraphs if p.startswith('Fig. ')]==[1,2,3]
    assert not re.search(r'\bmodules?\b',text,re.I)
    expected=['I. INTRODUCTION','II. LITERATURE SURVEY','III. PROPOSED SYSTEMS','IV. RESULTS AND IMPLEMENTATION','V. PERFORMANCE AND METRICS','VI. CONCLUSION','REFERENCES']
    assert [p.text for p in doc.paragraphs if p.style.name=='Heading 1']==expected
    diagram_index=next(i for i,p in enumerate(paragraphs) if p.startswith('Fig. 1.'))
    assert paragraphs.index('III. PROPOSED SYSTEMS')<diagram_index<paragraphs.index('A. Preprocessing and Validation')
    assert len(doc.element.xpath('.//m:oMath'))==6
    assert doc.styles['Normal'].font.size.pt==10
    assert doc.styles['Title'].font.size.pt==24
    assert doc.styles['Normal'].font.name=='Times New Roman'
    assert any(s._sectPr.find(qn('w:cols')).get(qn('w:num'))=='2' for s in doc.sections)
    assert 'Tanay Mohta and Aditya Rawat' in text
    assert not any(token in text for token in ['{{','TBD','YOUR NAME','uploaded'])
    metrics=load(RUN/'semeval-development/metrics.csv')
    table_text='\n'.join(c.text for t in doc.tables for r in t.rows for c in r.cells)
    for row in metrics:
        assert f'{100*float(row["accuracy"]):.2f}' in table_text
        for key in ['macro_f1','tuple_precision','tuple_recall','tuple_f1']:assert f'{float(row[key]):.4f}' in table_text
    for value in ['72/148','85/148']:assert value in table_text.replace(' ','')
    for row in load(RUN/'semeval-development/per_class_metrics.csv'):
        if float(row['f1']):assert f'{float(row["f1"]):.4f}' in text
    dl_metrics=load(DL/'metrics.csv')
    dl_table='\n'.join(c.text for r in doc.tables[3].rows for c in r.cells)
    for row in dl_metrics:
        for key in ('precision','recall','pair_micro_f1','exact_segment_match'):
            assert f'{100*float(row[key]):.2f}' in dl_table
        tp,ng,npred=map(int,(row['matched_pairs'],row['gold_pairs'],row['predicted_pairs']))
        assert abs(float(row['pair_micro_f1'])-2*tp/(ng+npred))<1e-12
    settings=json.loads((DL_RUN/'frozen_settings.json').read_text())
    assert settings['encoder_hash_before']!=settings['encoder_hash_after']
    assert settings['data_provenance']['ai_training_segments']==0
    assert settings['admitted_training_segments']==2038 and settings['training_pairs']==16304
    for marker in ('1,474','1,113','724','414','383','456','73','97 correct','loses 44','Facilities declines','Neutral/Mixed','no AI-labelled'):
        assert marker in text,marker
    assert 'E. Independent Deep Learning Reanalysis' in paragraphs
    assert all('Module' not in p for p in paragraphs)
    table_captions=[p.split('\n')[0] for p in paragraphs if p.startswith('TABLE ')]
    assert table_captions==['TABLE I','TABLE II','TABLE III','TABLE IV','TABLE V']
    category_table='\n'.join(c.text for r in doc.tables[-1].rows for c in r.cells)
    for r in load(DL/'per_category_comparison.csv'):assert f'{float(r["pair_micro_f1"]):.4f}' in category_table
    architecture=(OUT/'assets/architecture_approved.svg').read_text(encoding='utf-8')
    for marker in ('SemEval','Booking','3,089','500','Trip-MA','2,038','602','1,474','Evaluation Only'):
        assert marker in architecture,marker
    for removed in ('Retain Hybrid','Failure Record','Hybrid status?'):
        assert removed not in architecture
    # Preserve the original IEEE body typography, page size and margins.
    from io import BytesIO
    with ZipFile(OUT/'archive/pre-eight-page-revision-2026-10-26.zip') as z:
        prior=Document(BytesIO(z.read('paper/Team_17_IEEE_Conference_Paper.docx')))
    for attribute in ('page_width','page_height','top_margin','bottom_margin','left_margin','right_margin'):
        assert getattr(doc.sections[0],attribute)==getattr(prior.sections[0],attribute),attribute
    for attribute in ('size','name'):assert getattr(doc.styles['Normal'].font,attribute)==getattr(prior.styles['Normal'].font,attribute)
    assert not any(token in body for token in ['28.38%','0.2260','42/148','expanded hybrid','frozen baseline','earlier hybrid'])
    assert '21,181' in body and '14,311' in body and '772' in body
    coverage_evidence={}
    for cohort in ['booking-cleaned500','booking-all']:
        coverage_evidence[cohort]=load(RUN/cohort/'coverage.csv')
        for row in coverage_evidence[cohort]:
            if row['status'] in {'complete','partial','unsupported'}:assert row['reviews'] in table_text
    with pdfplumber.open(OUT/'Team_17_IEEE_Conference_Paper.pdf') as pdf:
        pages=len(pdf.pages);assert pages==8,pages
        fonts=sorted({c['fontname'] for p in pdf.pages for c in p.chars})
        assert all('TimesNewRoman' in f or 'CambriaMath' in f for f in fonts),fonts
        assert all(not any(c['x0']<32 or c['x1']>565 or c['top']<45 or c['bottom']>725 for c in p.chars) for p in pdf.pages)
    page_dir=ROOT/'tmp/paper-qa-eight-pages-2026-10-27'
    assert len(list(page_dir.glob('page-*.png')))==pages
    from verify_paper_graph_provenance import verify
    provenance=verify()
    audit={'updated':'2026-10-27','pages':pages,'research_works':32,'total_references':35,'pseudocode_algorithms':5,
           'architecture_diagrams':1,'measured_charts':2,'tables':5,'native_math_equations':6,
           'architecture_revision':'User-approved five-stage architecture spans both columns; datasets use connected data blocks, training and evaluation are separate, and the detailed status-decision section is removed. Only Partial/Unsupported reviews feed DL.',
           'compression':'Repeated prose and redundant figures removed; original IEEE typography, margins and page size unchanged.',
           'algorithm_revision':'Five compact pseudocodes; Algorithm 5 implements only Partial/Unsupported whole-review reanalysis with the fine-tuned DL checkpoint.',
           'algorithm_style_sources':[
               'https://journals.ieeeauthorcenter.ieee.org/wp-content/uploads/sites/7/Editing-Mathematics.pdf',
               'https://ctan.math.washington.edu/tex-archive/macros/latex2e/contrib/algorithmicx/algorithmicx.pdf'],
           'reference_order':'all 35 cited in first-appearance order',
           'metric_table':'SemEval lexical and Trip-MA DL tables match separate saved metric CSVs',
           'author_details':'two supplied names; affiliations/emails intentionally omitted until supplied',
           'evaluation_scope':'development/exploratory; not untouched holdout',
           'rendering':'Microsoft Word 16 native export plus canonical DOCX skill rasterizer',
           'visual_review':f'all {pages} final rendered pages inspected; final reference page balanced',
           'experiment_root':str(RUN.relative_to(ROOT)),
           'comparison':'SentiWordNet versus Hybrid on mapped SemEval; Hybrid versus Hybrid-DL on mapped Trip-MA; no cross-dataset ranking',
           'dl_checkpoint':json.loads((DL/'evaluation_manifest.json').read_text())['model_id'],
           'dl_training_settings':str((DL_RUN/'frozen_settings.json').relative_to(ROOT)),
           'dl_results':dl_metrics,
           'booking_cleaning_manifest':'output/datasets/booking_cleaned/cleaning_manifest.json',
           'results':{r['system']:{key:r[key] for key in ['cases','correct','accuracy','macro_f1','tuple_f1']} for r in metrics},
           'prior_version':'archive/pre-approved-dataset-diagram-2026-10-27.zip',
           'fonts':fonts,
           'sha256':{p.name:sha256(p.read_bytes()).hexdigest() for p in [path,OUT/'Team_17_IEEE_Conference_Paper.pdf']}}
    (OUT/'paper_audit.json').write_text(json.dumps(audit,indent=2),encoding='utf-8')
    with ZipFile(OUT/'Team_17_Paper_Figures_and_Tables.zip','w',compression=ZIP_DEFLATED) as archive:
        for stem in FIGURES:
            for suffix in ['.svg','.pdf','.png']:
                asset=OUT/'assets'/(stem+suffix);archive.write(asset,'figures/'+asset.name)
        for directory,names in {
            'semeval-development':['metrics.csv','per_class_metrics.csv','dataset_description.csv','annotation_scope.csv','excluded_annotations.csv'],
            'random200':['coverage.csv','dataset_description.csv','module_results.csv'],
            'booking-cleaned500':['coverage.csv','dataset_description.csv','module_results.csv'],
            'booking-all':['coverage.csv'],
        }.items():
            for name in names:archive.write(RUN/directory/name,f'tables/{directory}/{name}')
        archive.write(OUT/'references.json','references.json')
        archive.write(OUT/'PAPER_README.md','README.md')
        archive.write(OUT/'paper_audit.json','paper_audit.json')
        archive.write(OUT/'graph_provenance.json','graph_provenance.json')
        archive.write(ROOT/'docs/paper-graph-provenance.md','graph_provenance.md')
        archive.write(ROOT/'output/datasets/booking_cleaned/cleaning_manifest.json','tables/booking-cleaning/cleaning_manifest.json')
        for name in ['metrics.csv','per_category_comparison.csv','taxonomy_sensitivity.csv','evaluation_manifest.json','change_summary.json']:
            archive.write(DL/name,'tables/tripma-dl/'+name)
        for name in ['epoch_metrics.csv','frozen_settings.json']:
            archive.write(DL_RUN/name,'tables/dl-training/'+name)
        archive.write(ROOT/'output/neural/2026-10-23-dl-pair-booking500/summary.json','tables/dl-booking500/summary.json')
    print(json.dumps({k:v for k,v in audit.items() if k not in {'fonts','sha256'}},indent=2))


if __name__=='__main__':main()
