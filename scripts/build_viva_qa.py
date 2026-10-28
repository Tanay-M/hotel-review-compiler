"""Maintain the user's growing research viva question-and-answer document."""
from pathlib import Path
import json
from docx import Document
from docx.shared import Inches, Pt, RGBColor

ROOT=Path(__file__).resolve().parents[1]
DATA=ROOT/'docs/research-paper-viva-qa.json'
OUT=ROOT/'Research_Paper_Viva_QA.docx'

INITIAL=[
 {'id':1,'topic':'Abstract','question':'What is selective deep learning reanalysis?',
  'points':[
   'Selective means only Hybrid reviews marked Partial or Unsupported enter the deep learning stage. Complete reviews retain Hybrid; invalid inputs and component failures are not silently reanalysed.',
   'Reanalysis means the model independently reads every original sentence in the routed review. It does not receive only unmatched words, and it does not append predictions to the Hybrid opinion list.',
   'The DL result replaces the selected analysis for that routed review, including abstention when no category is accepted. Original Hybrid evidence remains available.',
   'Complete is a processing status, not proof that the review is correct. Selective routing can improve some cases and harm others.'],
  'viva':'We apply our fine-tuned deep learning model only to reviews the Hybrid system marks Partial or Unsupported. It rereads their original text independently and produces category-sentiment results or abstains; Complete Hybrid reviews are retained.'},
 {'id':2,'topic':'Abstract','question':'What are a MiniLM classifier, encoder and neural head?',
  'points':[
   'MiniLM is a compact pretrained transformer language model. In this project we initialize from all-MiniLM-L6-v2 and fine-tune it for classification; it is not a generative chatbot.',
   'Input is a sentence paired with a category description, for example a review sentence plus the Service query. The tokenizer converts the pair into numeric tokens.',
   'The encoder learns contextual numerical representations of the paired text. The classification head turns that representation into scores for Absent, Negative, Neutral/Mixed and Positive.',
   'Softmax converts output scores into probabilities summing to one. An accepted category must satisfy its validation-selected threshold; probability is not guaranteed correctness or sentiment intensity.',
   'The same encoder and head are reused across eight queries: five hotel categories and three background groups. Staff is pooled into Service in the DL taxonomy. Both encoder and head were fine-tuned on human-labelled Trip-MA data.'],
  'viva':'The MiniLM encoder converts the review sentence and category query into a contextual representation. A small neural classification head predicts whether that category is absent or has negative, neutral/mixed or positive sentiment. We fine-tuned both parts for hotel reviews.'},
 {'id':3,'topic':'Abstract','question':'What is polarity accuracy?',
  'points':[
   'Polarity means sentiment direction: Positive, Negative or Neutral in the SemEval classification comparison.',
   'Polarity accuracy = correct single-polarity predictions / all eligible sentence-category cases x 100. The denominator is 148 cases, not 148 whole reviews.',
   'SentiWordNet: 72/148 x 100 = 48.65%. Hybrid: 85/148 x 100 = 57.43%. Missing or Mixed predictions count as errors because the selected gold cases expect one label.',
   'Strong/Weak Positive collapse to Positive and Strong/Weak Negative to Negative for this three-label metric. We do not claim independently evaluated five-band accuracy.'],
  'viva':'Polarity accuracy measures how often the predicted sentiment direction matches the human annotation. On 148 eligible sentence-category cases, SentiWordNet correctly labels 72 and Hybrid correctly labels 85, giving 48.65% and 57.43%.'},
 {'id':4,'topic':'Abstract','question':'Explain the SemEval pair-F1 and Trip-MA deep learning improvement statement.',
  'points':[
   'A category-polarity pair is an item such as (Room, Positive) or (Service, Negative), associated with a particular sentence or segment. Both category and polarity must match the gold annotation.',
   'Precision = TP/(TP+FP); recall = TP/(TP+FN); F1 = 2PR/(P+R) = 2TP/(2TP+FP+FN). TP is a correct predicted pair, FP an extra or wrong pair, and FN a missing gold pair.',
   'SemEval pair evaluation has 162 gold pairs. SentiWordNet matches 81 of 112 predictions: F1 = 2 x 81/(112+162) = 0.5912. Hybrid matches 99 of 145 predictions: F1 = 2 x 99/(145+162) = 0.6450. Its 148-case polarity accuracy is a separate calculation.',
   'Trip-MA is a separate mapped experiment on 1,474 annotated test segments. Across all segments, Hybrid matches 314 pairs and the selective pipeline matches 367: micro-F1 rises from 0.4184 to 0.4993, an increase of 0.0809 or 8.09 F1 percentage points.',
   'Of those segments, 1,113 are Partial/Unsupported and enter DL; the other 361 outputs remain unchanged. Within the routed subset, correct pairs rise from 127 to 180 and F1 rises from 0.3261 to 0.4813: 15.52 F1 percentage points.',
   'Micro-F1 pools pair counts before calculation; it is not average review F1 or the percentage of correctly analysed reviews. F1 of 0.4993 does not mean 49.93% review accuracy.',
   'The cohorts and category mappings differ, so SemEval SentiWordNet scores cannot be directly ranked against Trip-MA DL scores. Results show recovery in this exploratory evaluation, not guaranteed correction of every Partial review; Facilities F1 declines.'],
  'viva':'On SemEval, Hybrid improves sentiment-label accuracy and category-polarity pair recovery over SentiWordNet. Separately, on Trip-MA, sending only unresolved Hybrid segments to our fine-tuned model raises overall pair-F1 from 0.4184 to 0.4993; within the routed subset it rises from 0.3261 to 0.4813. These are F1 improvements, not whole-review accuracy percentages.'}
]


def build():
    if not DATA.exists():
        DATA.write_text(json.dumps(INITIAL,indent=2,ensure_ascii=False),encoding='utf-8')
    entries=json.loads(DATA.read_text(encoding='utf-8'))
    doc=Document();s=doc.sections[0]
    s.page_width=Inches(8.5);s.page_height=Inches(11)
    s.top_margin=s.bottom_margin=Inches(.65)
    s.left_margin=s.right_margin=Inches(.75)
    for name in ('Normal','Title','Heading 1','Heading 2'):
        style=doc.styles[name];style.font.name='Calibri';style.font.color.rgb=RGBColor(0,0,0)
    doc.styles['Normal'].font.size=Pt(10.5)
    doc.styles['Normal'].paragraph_format.space_after=Pt(5)
    doc.styles['Normal'].paragraph_format.line_spacing=1.02
    doc.styles['Title'].font.size=Pt(21)
    doc.styles['Heading 1'].font.size=Pt(13)
    doc.styles['Heading 1'].paragraph_format.space_before=Pt(10)
    doc.add_paragraph('Research Paper Viva Questions and Answers','Title')
    doc.add_paragraph('Team 17 | Tanay Mohta and Aditya Rawat | Updated 27 October 2026')
    doc.add_paragraph('Study notes and short faculty-facing answers. New questions from this discussion will be added to this same document. The conference paper is unchanged.')
    for item in entries:
        if item['id']==3:doc.add_page_break()
        doc.add_paragraph(f"Q{item['id']}  {item['question']}",'Heading 1')
        doc.add_paragraph('Paper section: '+item['topic'])
        for point in item['points']:doc.add_paragraph(point,'List Bullet')
        p=doc.add_paragraph();p.add_run('Answer to faculty: ').bold=True;p.add_run(item['viva'])
    for node in list(doc.element.xpath('.//w:pBdr'))+list(doc.styles.element.xpath('.//w:pBdr')):
        node.getparent().remove(node)
    doc.save(OUT)
    print(f'Saved {len(entries)} questions to {OUT}')


if __name__=='__main__':build()
